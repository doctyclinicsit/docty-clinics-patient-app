from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_DOCUMENTATION.md"
OUTPUT = ROOT / "docs" / "Docty_Clinics_Project_Handbook.docx"
LOGO = ROOT / "public" / "docty-logo-full.png"

RED = "FE065C"
BLUE = "0BB8FC"
NAVY = "092F49"
INK = "183B50"
MUTED = "5D7180"
LIGHT_BLUE = "EAF8FE"
LIGHT_RED = "FFF0F5"
LIGHT_GRAY = "F4F7F9"
WHITE = "FFFFFF"
BORDER = "D7E3EA"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_font(run, size=9, color=MUTED)


def add_inline_markdown(paragraph, text, base_size=11, base_color=INK):
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|https?://\S+)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_font(paragraph.add_run(text[cursor:match.start()]), size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=9.5, color=NAVY)
            run.font.highlight_color = None
        elif token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=base_size, color=base_color, bold=True)
        elif token.startswith("*"):
            set_font(paragraph.add_run(token[1:-1]), size=base_size, color=base_color, italic=True)
        else:
            run = paragraph.add_run(token.rstrip(".,)"))
            set_font(run, size=base_size, color=BLUE)
            run.underline = True
            if token[-1:] in ".,)":
                set_font(paragraph.add_run(token[-1]), size=base_size, color=base_color)
        cursor = match.end()
    if cursor < len(text):
        set_font(paragraph.add_run(text[cursor:]), size=base_size, color=base_color)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in (
        ("Heading 1", 16, RED, 18, 8),
        ("Heading 2", 13, NAVY, 14, 7),
        ("Heading 3", 11.5, BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.18


def add_running_furniture(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("DOCTY CLINICS"), size=8.5, color=RED, bold=True)
    set_font(p.add_run("  |  Project and Operations Handbook"), size=8.5, color=MUTED)

    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), BORDER)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_p = footer.paragraphs[0]
    set_font(footer_p.add_run("Confidential operational reference  |  "), size=8.5, color=MUTED)
    add_page_number(footer_p)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO.exists():
        picture = p.add_run().add_picture(str(LOGO), width=Inches(2.35))
        picture._inline.docPr.set("descr", "Docty Clinics logo")
        picture._inline.docPr.set("title", "Docty Clinics")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(78)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("PROJECT AND OPERATIONS"), size=12, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Docty Clinics Patient Application"), size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(36)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Architecture, integrations, operations, security, and maintenance handbook"), size=13, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2200, 7160])
    set_repeat_table_header(table.rows[0])
    table.style = "Table Grid"
    metadata = [
        ("Version", "1.0"),
        ("Document date", "June 19, 2026"),
        ("Production", "docty-clinics-patient-app-eta.vercel.app"),
        ("Application", "docty-clinics-patient-app"),
    ]
    for index, (label, value) in enumerate(metadata):
        left, right = table.rows[index].cells
        shade_cell(left, LIGHT_BLUE if index % 2 == 0 else LIGHT_RED)
        shade_cell(right, WHITE)
        set_cell_border(left)
        set_cell_border(right)
        left_p = left.paragraphs[0]
        right_p = right.paragraphs[0]
        set_font(left_p.add_run(label), size=9.5, color=NAVY, bold=True)
        set_font(right_p.add_run(value), size=9.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(
        p.add_run("Do not store API keys, access tokens, signed URLs, or session secrets in documentation or source control."),
        size=9.5,
        color=RED,
        bold=True,
    )


def add_contents(doc, headings):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run("CONTENTS"), size=10, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    set_font(p.add_run("Handbook sections"), size=23, color=NAVY, bold=True)

    for title in headings:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(title), size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(0)
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), LIGHT_BLUE)
    p._p.get_or_add_pPr().append(shade)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    set_font(p.add_run("Reading guide: Sections 1-8 explain the product; 9-16 cover APIs, integrations, security, and deployment; 17-25 cover behavior, testing, maintenance, and change management."), size=9.5, color=NAVY)


def add_markdown_table(doc, rows):
    if not rows:
        return
    columns = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    if columns == 2:
        widths = [2600, 6760]
    elif columns == 3:
        widths = [2200, 2200, 4960]
    elif columns == 4:
        widths = [1500, 1900, 2300, 3660]
    else:
        base = CONTENT_WIDTH_DXA // columns
        widths = [base] * columns
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])

    for row_index, source_row in enumerate(rows):
        for col_index, value in enumerate(source_row):
            cell = table.rows[row_index].cells[col_index]
            shade_cell(cell, NAVY if row_index == 0 else (LIGHT_GRAY if row_index % 2 == 0 else WHITE))
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_markdown(
                p,
                value,
                base_size=9 if row_index == 0 else 8.8,
                base_color=WHITE if row_index == 0 else INK,
            )
            for run in p.runs:
                if row_index == 0:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_code_block(doc, code_lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, "102D40")
    set_cell_border(cell, color="102D40", size="2")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run("\n".join(code_lines)), name="Consolas", size=8.3, color=WHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_RED)
    set_cell_border(cell, color="FFC0D5")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("IMPORTANT  "), size=9.5, color=RED, bold=True)
    add_inline_markdown(p, text, base_size=9.5, base_color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def build_document():
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    section_headings = [
        line[3:].strip()
        for line in lines
        if line.startswith("## ") and re.match(r"\d+\.", line[3:].strip())
    ]

    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_contents(doc, section_headings)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.page_width = Inches(8.5)
    body_section.page_height = Inches(11)
    body_section.top_margin = Inches(0.8)
    body_section.bottom_margin = Inches(0.75)
    body_section.left_margin = Inches(1)
    body_section.right_margin = Inches(1)
    body_section.header_distance = Inches(0.35)
    body_section.footer_distance = Inches(0.35)
    body_section.different_first_page_header_footer = False
    add_running_furniture(body_section)

    start = next(
        index
        for index, line in enumerate(lines)
        if line.startswith("## 1. ")
    )
    index = start
    in_code = False
    code_lines = []

    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not stripped or stripped == "---":
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1)) - 1
            title = heading_match.group(2)
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_inline_markdown(p, title, base_size={1: 16, 2: 13, 3: 11.5}[min(level, 3)], base_color={1: RED, 2: NAVY, 3: BLUE}[min(level, 3)])
            index += 1
            continue

        if stripped.startswith("> "):
            add_note(doc, stripped[2:])
            index += 1
            continue

        bullet_match = re.match(r"^(\s*)-\s+(.+)$", line)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            nesting = len(bullet_match.group(1)) // 2
            p.paragraph_format.left_indent = Inches(0.375 + nesting * 0.22)
            add_inline_markdown(p, bullet_match.group(2))
            index += 1
            continue

        number_match = re.match(r"^(\s*)\d+\.\s+(.+)$", line)
        if number_match:
            p = doc.add_paragraph(style="List Number")
            nesting = len(number_match.group(1)) // 2
            p.paragraph_format.left_indent = Inches(0.375 + nesting * 0.22)
            add_inline_markdown(p, number_match.group(2))
            index += 1
            continue

        p = doc.add_paragraph()
        add_inline_markdown(p, stripped)
        index += 1

    doc.core_properties.title = "Docty Clinics Patient Application - Project and Operations Handbook"
    doc.core_properties.subject = "Architecture, integrations, deployment, security, testing, and maintenance"
    doc.core_properties.author = "Docty Clinics"
    doc.core_properties.keywords = "Docty Clinics, patient app, Eka Care, MSG91, Zoho CRM, Vercel"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
